"""Extract a small, read-only text preview from existing Office assets."""

from __future__ import annotations

import json
import base64
import sys
from pathlib import Path


def clean(value: object) -> str:
    return " ".join(str(value or "").split()).strip()


def limited(values: list[str], limit: int = 12) -> list[str]:
    out: list[str] = []
    for value in values:
        text = clean(value)
        if not text or text in out:
            continue
        out.append(text[:90])
        if len(out) >= limit:
            break
    return out


def extract_xlsx(path: Path) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        lines: list[str] = []
        for ws in wb.worksheets[:2]:
            lines.append(ws.title)
            for row in ws.iter_rows(max_row=min(ws.max_row or 1, 24), max_col=min(ws.max_column or 1, 10)):
                cells = [clean(cell.value) for cell in row if clean(cell.value)]
                if cells:
                    lines.append(" / ".join(cells))
                if len(lines) >= 12:
                    break
            if len(lines) >= 12:
                break
        return limited(lines)
    finally:
        wb.close()


def extract_docx(path: Path) -> list[str]:
    from docx import Document

    doc = Document(str(path))
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables[:3]:
        for row in table.rows[:12]:
            lines.append(" / ".join(clean(cell.text) for cell in row.cells if clean(cell.text)))
    return limited(lines)


def extract_pptx(path: Path) -> list[str]:
    from pptx import Presentation

    prs = Presentation(str(path))
    lines: list[str] = []
    for index, slide in enumerate(list(prs.slides)[:3], 1):
        slide_lines = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = clean(getattr(shape, "text", ""))
                if text:
                    slide_lines.append(text)
        if slide_lines:
            lines.append(f"Slide {index}: " + " / ".join(slide_lines))
    return limited(lines)


def main() -> None:
    payload = json.loads(base64.b64decode(sys.stdin.buffer.read()).decode("utf-8"))
    results = []
    for item in payload:
        row = {"id": item.get("id"), "ok": False, "lines": [], "reason": None}
        try:
            path = Path(item["path"])
            ext = path.suffix.lower()
            if ext == ".xlsx":
                lines = extract_xlsx(path)
            elif ext == ".docx":
                lines = extract_docx(path)
            elif ext == ".pptx":
                lines = extract_pptx(path)
            else:
                raise ValueError(f"unsupported_format:{ext}")
            if not lines:
                raise ValueError("empty_preview_content")
            row.update({"ok": True, "lines": lines})
        except Exception as exc:  # noqa: BLE001
            row["reason"] = str(exc)
        results.append(row)
    json.dump(results, sys.stdout, ensure_ascii=False)


if __name__ == "__main__":
    main()
